import pytest
from pathlib import Path
import os
import sys
from openpyxl import load_workbook

# Add parent dir to path so we can import pipeline
sys.path.append(str(Path(__file__).resolve().parent.parent))
import pipeline

ROOT = Path(__file__).resolve().parent.parent
CENSUS = ROOT / "project_census.xlsx"

def test_excel_integrity():
    """Verify project_census.xlsx has all required sheets and headers."""
    assert CENSUS.exists(), "Census file should exist"
    wb = load_workbook(CENSUS, read_only=True)

    assert "Mapping" in wb.sheetnames
    assert "Inventory" in wb.sheetnames
    assert "Enrichment" in wb.sheetnames

    mapping_ws = wb["Mapping"]
    headers = [cell.value for cell in mapping_ws[1]]
    assert "Sutta ID" in headers
    assert "Video ID" in headers
    assert "Nikaya Folder" in headers

def test_id_normalization():
    """Ensure various Sutta ID formats normalize to x_y_z."""
    assert pipeline.normalize_id("5.4.40") == "5_4_40"
    assert pipeline.normalize_id("5-4-40") == "5_4_40"
    assert pipeline.normalize_id("5_4_40") == "5_4_40"
    assert pipeline.normalize_id("5.1") == "5_1"

def test_path_sanitization():
    """Ensure Nikaya folder names are safe for filesystem."""
    unsafe = "Anguttara Nikaya: Book 5 / Chapter 4"
    safe = pipeline.safe_name(unsafe)
    assert ":" not in safe
    assert "/" not in safe
    assert safe == "Anguttara Nikaya_ Book 5 _ Chapter 4"

def test_sha256_reproducibility(tmp_path):
    """Ensure hashing logic is consistent."""
    f = tmp_path / "test.txt"
    f.write_text("buddha", encoding="utf-8")
    h1 = pipeline.sha256(f)
    h2 = pipeline.sha256(f)
    assert h1 == h2
    assert len(h1) == 64
